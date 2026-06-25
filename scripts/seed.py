"""Idempotent demo seed: users, tool registry, role grants.

Demo credentials (dev only):
  admin@example.com / admin123     role=admin
  analyst@example.com / analyst123 role=analyst
  viewer@example.com / viewer123   role=viewer
"""

import asyncio
from pathlib import Path

import httpx
from sqlalchemy import select

from app.auth import hash_password
from app.config import settings
from app.db import SessionLocal
from app.models import RoleToolGrant, Tool, User
from app.services.ingest import ingest_directory

DATA_DIR = Path(__file__).resolve().parent.parent / "data"
CORPUS_DIR = DATA_DIR / "corpus"

USERS = [
    ("admin@example.com", "admin123", "admin"),
    ("analyst@example.com", "analyst123", "analyst"),
    ("viewer@example.com", "viewer123", "viewer"),
]

TOOLS = [
    ("search_docs", "docs", "Semantic search over the internal documentation corpus", "read"),
    ("list_tables", "db", "List queryable tables in the analytics database", "read"),
    ("run_query", "db", "Run a read-only SQL SELECT against the analytics database", "read"),
    ("create_ticket", "tickets",
     "Create a ticket in the ticketing system (requires human approval)", "critical"),
]

GRANTS = {
    "viewer": ["search_docs"],
    "analyst": ["search_docs", "list_tables", "run_query", "create_ticket"],
    "admin": ["search_docs", "list_tables", "run_query", "create_ticket"],
}


def load_retail() -> None:
    """Bulk-load the UCI Online Retail xlsx into sales.transactions (COPY, idempotent)."""
    import psycopg
    from openpyxl import load_workbook

    xlsx = DATA_DIR / "online_retail.xlsx"
    if not xlsx.exists():
        print("skipped retail load — run `make data` first")
        return
    sync_url = settings.database_url.replace("+asyncpg", "")
    with psycopg.connect(sync_url) as conn:
        if conn.execute("select count(*) from sales.transactions").fetchone()[0]:
            return
        print("loading UCI Online Retail into sales.transactions (~540k rows)...")
        sheet = load_workbook(xlsx, read_only=True).active
        with conn.cursor() as cur, cur.copy(
            "copy sales.transactions (invoice_no, stock_code, description, quantity, "
            "invoice_date, unit_price, customer_id, country) from stdin"
        ) as copy:
            for row in sheet.iter_rows(min_row=2, values_only=True):
                invoice, stock, desc, qty, ts, price, customer, country = row[:8]
                if invoice is None or ts is None:
                    continue
                copy.write_row((str(invoice), str(stock), desc, int(qty), ts, price,
                                int(customer) if customer else None, country or "Unknown"))
        conn.commit()
        n = conn.execute("select count(*) from sales.transactions").fetchone()[0]
        print(f"loaded {n} transactions")


async def seed() -> None:
    async with SessionLocal() as session:
        for email, password, role in USERS:
            exists = (await session.execute(select(User).where(User.email == email))).scalar_one_or_none()
            if not exists:
                session.add(User(email=email, password_hash=hash_password(password), role=role))
        for name, server, description, risk in TOOLS:
            if await session.get(Tool, name) is None:
                session.add(Tool(name=name, mcp_server=server, description=description, risk=risk))
        await session.flush()
        for role, tool_names in GRANTS.items():
            for tool_name in tool_names:
                if await session.get(RoleToolGrant, (role, tool_name)) is None:
                    session.add(RoleToolGrant(role=role, tool_name=tool_name))
        await session.commit()
    print("seeded users, tools and grants")

    await asyncio.to_thread(load_retail)

    if not CORPUS_DIR.exists():
        print("skipped corpus ingest — run `make data` first")
        return
    try:
        httpx.get(f"{settings.ollama_base_url}/api/version", timeout=2)
        chunks = await ingest_directory(CORPUS_DIR)
        print(f"ingested corpus: {chunks} chunks")
    except Exception as exc:  # corpus is optional at seed time
        print(f"skipped corpus ingest ({type(exc).__name__}: {exc}) — run again once ollama is up")


if __name__ == "__main__":
    asyncio.run(seed())
